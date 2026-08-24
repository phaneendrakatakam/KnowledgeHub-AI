import sys
from pathlib import Path

import pytest
from docx import Document

sys.path.insert(
    0,
    str(
        Path(__file__).resolve().parents[1]
        / "backend"
    )
)

from parsers import get_parser, parse_document
from parsers.docx_parser import DOCXParser
from parsers.markdown_parser import MarkdownParser
from parsers.txt_parser import TextParser


def test_get_parser_returns_text_parser():
    assert isinstance(get_parser("notes.txt"), TextParser)


def test_get_parser_returns_markdown_parser():
    assert isinstance(get_parser("README.md"), MarkdownParser)


def test_get_parser_supports_markdown_extension():
    assert isinstance(get_parser("guide.markdown"), MarkdownParser)


def test_get_parser_returns_docx_parser():
    assert isinstance(get_parser("runbook.docx"), DOCXParser)


def test_get_parser_rejects_unsupported_file():
    with pytest.raises(ValueError, match="Unsupported file type"):
        get_parser("document.exe")


def test_text_parser_reads_file(tmp_path):
    file_path = tmp_path / "notes.txt"
    file_path.write_text(
        "KnowledgeHub V3 supports text files.",
        encoding="utf-8",
    )

    results = parse_document(str(file_path))
    assert len(results) == 1

    result = results[0]
    assert result.content == "KnowledgeHub V3 supports text files."
    assert result.filename == "notes.txt"
    assert result.file_type == "txt"
    assert result.page_number is None
    assert result.section is None
    assert result.sheet_name is None
    assert result.row_start is None
    assert result.row_end is None
    assert result.metadata == {"source_type": "text"}


def test_markdown_parser_reads_file(tmp_path):
    file_path = tmp_path / "guide.md"
    file_path.write_text(
        "# KnowledgeHub V3\n\nMulti-format ingestion is enabled.",
        encoding="utf-8",
    )

    results = parse_document(str(file_path))
    assert len(results) == 1

    result = results[0]
    assert result.content == (
        "# KnowledgeHub V3\n\nMulti-format ingestion is enabled."
    )
    assert result.filename == "guide.md"
    assert result.file_type == "markdown"
    assert result.metadata == {"source_type": "markdown"}


def test_docx_parser_reads_paragraph_content(tmp_path):
    file_path = tmp_path / "runbook.docx"

    document = Document()
    document.add_paragraph("KnowledgeHub V3 supports DOCX files.")
    document.save(file_path)

    results = parse_document(str(file_path))

    assert len(results) == 1
    assert results[0].content == "KnowledgeHub V3 supports DOCX files."
    assert results[0].filename == "runbook.docx"
    assert results[0].file_type == "docx"
    assert results[0].section is None
    assert results[0].metadata == {"source_type": "docx"}


def test_docx_parser_preserves_heading_as_section(tmp_path):
    file_path = tmp_path / "architecture.docx"

    document = Document()
    document.add_heading("Deployment", level=1)
    document.add_paragraph("The application is deployed after validation.")
    document.add_heading("Rollback", level=1)
    document.add_paragraph("Rollback restores a known-good version.")
    document.save(file_path)

    results = parse_document(str(file_path))

    assert len(results) == 2
    assert results[0].section == "Deployment"
    assert results[0].content == "The application is deployed after validation."
    assert results[1].section == "Rollback"
    assert results[1].content == "Rollback restores a known-good version."


def test_docx_parser_extracts_table_content(tmp_path):
    file_path = tmp_path / "deployments.docx"

    document = Document()
    document.add_heading("Deployment Matrix", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Environment"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Production"
    table.cell(1, 1).text = "Healthy"
    document.save(file_path)

    results = parse_document(str(file_path))

    assert len(results) == 1
    assert results[0].section == "Deployment Matrix"
    assert "Environment | Status" in results[0].content
    assert "Production | Healthy" in results[0].content


def test_text_parser_rejects_empty_file(tmp_path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("   \n\n   ", encoding="utf-8")

    with pytest.raises(
        ValueError,
        match="does not contain any readable content",
    ):
        TextParser().parse(str(file_path))


def test_markdown_parser_rejects_empty_file(tmp_path):
    file_path = tmp_path / "empty.md"
    file_path.write_text("\n\n", encoding="utf-8")

    with pytest.raises(
        ValueError,
        match="does not contain any readable content",
    ):
        MarkdownParser().parse(str(file_path))


def test_docx_parser_rejects_empty_file(tmp_path):
    file_path = tmp_path / "empty.docx"
    Document().save(file_path)

    with pytest.raises(
        ValueError,
        match="does not contain any readable content",
    ):
        DOCXParser().parse(str(file_path))


def test_text_parser_rejects_missing_file(tmp_path):
    missing_file = tmp_path / "missing.txt"

    with pytest.raises(FileNotFoundError, match="File not found"):
        TextParser().parse(str(missing_file))


def test_markdown_parser_rejects_missing_file(tmp_path):
    missing_file = tmp_path / "missing.md"

    with pytest.raises(FileNotFoundError, match="File not found"):
        MarkdownParser().parse(str(missing_file))


def test_docx_parser_rejects_missing_file(tmp_path):
    missing_file = tmp_path / "missing.docx"

    with pytest.raises(FileNotFoundError, match="File not found"):
        DOCXParser().parse(str(missing_file))